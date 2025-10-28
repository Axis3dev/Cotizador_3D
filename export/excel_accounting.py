"""Excel export utilities for accounting summaries."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Dict, Iterable

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from .branding import add_logo_to_sheet, resolve_logo_path

def export_accounting_excel(
    rows: Iterable[Dict[str, float]],
    filters: Dict[str, str],
    identity: Dict[str, object] | None = None,
    directory: Path | None = None,
    file_name: str = "reporte_contable.xlsx",
) -> Path:
    directory = directory or Path.cwd()
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / file_name

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"

    current_row = 1
    if identity:
        if add_logo_to_sheet(ws, identity, cell="A1"):
            current_row = 6
        elif resolve_logo_path(identity) is None:
            logging.getLogger(__name__).info("Logo no disponible para el Excel contable")
    ws.cell(row=current_row, column=1, value="Filtros")
    current_row += 1
    for key, value in filters.items():
        ws.cell(row=current_row, column=1, value=key)
        ws.cell(row=current_row, column=2, value=value)
        current_row += 1

    current_row += 1
    headers = ["Periodo", "Ingresos", "Costos", "Gastos", "Ganancia", "IVA"]
    ws.append(headers)
    for row in rows:
        ws.append(
            [
                row.get("periodo", ""),
                row.get("ingresos", 0.0),
                row.get("costos", 0.0),
                row.get("gastos", 0.0),
                row.get("ganancia", 0.0),
                row.get("iva", 0.0),
            ]
        )

    for idx, column in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(idx)].width = 18

    wb.save(path)
    return path


__all__ = ["export_accounting_excel"]
